import csv
import os
from html import escape
from pathlib import Path
from database.database import (
    get_connection,
    get_attendance_analytics,
    get_student_attendance,
    get_student_profile,
    get_student_academic_records,
)

class ReportGenerator:
    """
    Generates CSV and HTML/PDF formatted reports for EduPilot.
    """

    @staticmethod
    def _excel_workbook_type():
        try:
            from openpyxl import Workbook
        except ImportError as error:
            raise RuntimeError("Excel export needs openpyxl. Install the project dependencies from requirements.txt and try again.") from error
        return Workbook

    @staticmethod
    def export_attendance_csv(filepath: str, subject_id: int = None, start_date: str = None,
                              end_date: str = None, section: str = None,
                              class_id: int = None) -> bool:
        """
        Exports all attendance records to a CSV file.
        """
        conn = get_connection()
        c = conn.cursor()
        
        query = """
            SELECT 
                ats.class_date,
                s.code AS subject_code,
                s.name AS subject_name,
                st.roll_number,
                st.name AS student_name,
                COALESCE(ar.status, 'Absent') AS status,
                COALESCE(ar.verification_method, 'not recorded') AS verification_method,
                COALESCE(ar.marked_at, '') AS marked_at
            FROM attendance_sessions ats
            JOIN subjects s ON s.id = ats.subject_id
            JOIN students st ON (ats.class_id = st.class_id OR (ats.class_id IS NULL AND st.section = ats.section))
            LEFT JOIN attendance_records ar ON ar.session_id = ats.id AND ar.student_id = st.id
        """
        conditions, params = [], []
        if subject_id:
            conditions.append("s.id = ?")
            params.append(subject_id)
        if start_date:
            conditions.append("ats.class_date >= ?")
            params.append(start_date)
        if end_date:
            conditions.append("ats.class_date <= ?")
            params.append(end_date)
        if section:
            conditions.append("ats.section = ?")
            params.append(section)
        if class_id:
            conditions.append("st.class_id = ?")
            params.append(class_id)
        if conditions:
            query += " WHERE " + " AND ".join(conditions)

        query += " ORDER BY ats.class_date DESC, st.roll_number ASC"
        
        c.execute(query, params)
        rows = c.fetchall()
        conn.close()

        os.makedirs(os.path.dirname(os.path.abspath(filepath)), exist_ok=True)

        with open(filepath, mode="w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerow(["Class Date", "Subject Code", "Subject Name", "Roll Number", "Student Name", "Status", "Verification Method", "Marked At"])
            writer.writerows(rows)
            
        return True

    @staticmethod
    def export_attendance_xlsx(filepath: str, subject_id: int = None, start_date: str = None,
                               end_date: str = None, section: str = None,
                               class_id: int = None) -> bool:
        Workbook = ReportGenerator._excel_workbook_type()
        conn = get_connection()
        query = """SELECT ats.class_date, s.code, s.name, st.roll_number, st.name,
                COALESCE(ar.status, 'Absent'), COALESCE(ar.verification_method, 'not recorded'),
                COALESCE(ar.marked_at, '')
            FROM attendance_sessions ats JOIN subjects s ON s.id = ats.subject_id
            JOIN students st ON (ats.class_id = st.class_id OR (ats.class_id IS NULL AND st.section = ats.section))
            LEFT JOIN attendance_records ar ON ar.session_id = ats.id AND ar.student_id = st.id"""
        conditions, params = [], []
        if subject_id:
            conditions.append("s.id = ?")
            params.append(subject_id)
        if start_date:
            conditions.append("ats.class_date >= ?")
            params.append(start_date)
        if end_date:
            conditions.append("ats.class_date <= ?")
            params.append(end_date)
        if section:
            conditions.append("ats.section = ?")
            params.append(section)
        if class_id:
            conditions.append("st.class_id = ?")
            params.append(class_id)
        if conditions:
            query += " WHERE " + " AND ".join(conditions)
        rows = conn.execute(query + " ORDER BY ats.class_date DESC, st.roll_number", params).fetchall()
        conn.close()
        os.makedirs(os.path.dirname(os.path.abspath(filepath)), exist_ok=True)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Attendance"
        sheet.append(["Class Date", "Subject Code", "Subject Name", "Roll Number", "Student Name", "Status", "Verification Method", "Marked At"])
        for row in rows:
            sheet.append([ReportGenerator._excel_safe(value) for value in row])
        ReportGenerator._format_sheet(sheet)
        workbook.save(filepath)
        return True

    @staticmethod
    def export_attendance_pdf_report(filepath: str, title: str = "EduPilot Attendance Summary",
                                     start_date: str = None, end_date: str = None,
                                     subject_id: int = None, section: str = None,
                                     class_id: int = None) -> bool:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_LEFT
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

        rows = get_attendance_analytics(start_date, end_date, subject_id, section, class_id)
        os.makedirs(os.path.dirname(os.path.abspath(filepath)), exist_ok=True)

        if start_date and end_date and start_date == end_date:
            period_description = f"Date: {start_date}"
        elif start_date or end_date:
            period_description = f"Period: {start_date or 'Beginning'} - {end_date or 'Today'}"
        else:
            period_description = "Period: All available dates"
        if section:
            period_description += f" | Section {section}"

        sample = getSampleStyleSheet()
        title_style = ParagraphStyle(
            "AttendanceTitle", parent=sample["Title"], fontName="Helvetica-Bold",
            fontSize=20, leading=25, alignment=TA_LEFT, textColor=colors.HexColor("#263746"),
            spaceAfter=4,
        )
        subtitle_style = ParagraphStyle(
            "AttendanceSubtitle", parent=sample["Normal"], fontName="Helvetica",
            fontSize=11, leading=15, textColor=colors.HexColor("#52616B"),
        )
        cell_style = ParagraphStyle(
            "AttendanceCell", parent=sample["Normal"], fontName="Helvetica",
            fontSize=11, leading=14, textColor=colors.HexColor("#26323B"),
        )
        story = [
            Paragraph(escape(str(title)), title_style),
            Paragraph("Velammal College of Engineering &amp; Technology (VCET)", subtitle_style),
            Paragraph(escape(period_description), subtitle_style),
            Spacer(1, 7 * mm),
        ]
        table_data = [[
            "Roll Number", "Student Name", "Total Classes", "Classes Attended",
            "Attendance %", "Status",
        ]]
        status_values = []
        for roll, name, total, present, percentage in rows:
            status = "No sessions" if not total else ("Good" if percentage >= 75 else "Low (<75%)")
            status_values.append(status)
            table_data.append([
                Paragraph(escape(str(roll)), cell_style),
                Paragraph(escape(str(name)), cell_style),
                str(total), str(present),
                f"{percentage:.1f}%" if total else "-",
                status,
            ])

        page_width, _page_height = landscape(A4)
        available_width = page_width - 24 * mm
        column_weights = (1.15, 2.5, 1.25, 1.5, 1.2, 1.3)
        total_weight = sum(column_weights)
        column_widths = [available_width * weight / total_weight for weight in column_weights]
        table = Table(table_data, colWidths=column_widths, repeatRows=1, hAlign="LEFT")
        table_commands = [
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#26323B")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, 0), 10),
            ("LEADING", (0, 0), (-1, 0), 13),
            ("ALIGN", (0, 0), (-1, 0), "LEFT"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 8),
            ("RIGHTPADDING", (0, 0), (-1, -1), 8),
            ("TOPPADDING", (0, 0), (-1, -1), 8),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
            ("LINEBELOW", (0, 0), (-1, 0), 1, colors.HexColor("#48677D")),
            ("LINEBELOW", (0, 1), (-1, -1), 0.5, colors.HexColor("#DCE2E7")),
            ("ALIGN", (2, 1), (4, -1), "CENTER"),
            ("ALIGN", (5, 1), (5, -1), "CENTER"),
        ]
        for row_index in range(1, len(table_data)):
            if row_index % 2 == 0:
                table_commands.append(("BACKGROUND", (0, row_index), (-1, row_index), colors.HexColor("#F4F6F8")))
            status_color = "#48677D" if status_values[row_index - 1] == "Good" else "#7A6B52"
            table_commands.append(("TEXTCOLOR", (5, row_index), (5, row_index), colors.HexColor(status_color)))
        table.setStyle(TableStyle(table_commands))
        story.append(table)

        def draw_footer(canvas, document):
            canvas.saveState()
            canvas.setFont("Helvetica", 9)
            canvas.setFillColor(colors.HexColor("#52616B"))
            canvas.drawCentredString(page_width / 2, 8 * mm, f"EduPilot Attendance Report  |  Page {document.page}")
            canvas.restoreState()

        document = SimpleDocTemplate(
            filepath, pagesize=landscape(A4),
            leftMargin=12 * mm, rightMargin=12 * mm,
            topMargin=12 * mm, bottomMargin=17 * mm,
            title=str(title), author="EduPilot",
        )
        document.build(story, onFirstPage=draw_footer, onLaterPages=draw_footer)
        return os.path.isfile(filepath) and os.path.getsize(filepath) > 0

    @staticmethod
    def export_student_transcript_xlsx(user_id: int, filepath: str) -> bool:
        Workbook = ReportGenerator._excel_workbook_type()
        profile = get_student_profile(user_id)
        if not profile:
            return False
        os.makedirs(os.path.dirname(os.path.abspath(filepath)), exist_ok=True)
        workbook = Workbook()
        attendance = workbook.active
        attendance.title = "Attendance"
        attendance.append(["Subject Code", "Subject Name", "Sessions", "Present", "Attendance %"])
        for code, name, total, present in get_student_attendance(user_id):
            percent = round(present / total * 100, 1) if total else None
            attendance.append([ReportGenerator._excel_safe(code), ReportGenerator._excel_safe(name), total, present, percent])
        academics = workbook.create_sheet("Academic Marks")
        academics.append(["Subject Code", "Subject Name", "Internal /100", "Assignment /100", "Lab /100", "Total /300"])
        for row in get_student_academic_records(user_id):
            academics.append([ReportGenerator._excel_safe(value) if isinstance(value, str) else value for value in row])
        profile_sheet = workbook.create_sheet("Student")
        profile_sheet.append(["Name", "Roll Number", "Department", "Year", "Section"])
        profile_sheet.append([ReportGenerator._excel_safe(profile[3]), ReportGenerator._excel_safe(profile[2]),
                              ReportGenerator._excel_safe(profile[4]), profile[5], ReportGenerator._excel_safe(profile[6])])
        for sheet in workbook.worksheets:
            ReportGenerator._format_sheet(sheet)
        workbook.save(filepath)
        return True

    @staticmethod
    def _excel_safe(value):
        if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
            return "'" + value
        return value

    @staticmethod
    def _format_sheet(sheet):
        from openpyxl.styles import Font, PatternFill
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="354E60")
        for column in sheet.columns:
            width = min(48, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
            sheet.column_dimensions[column[0].column_letter].width = width

    @staticmethod
    def export_classroom_analytics_csv(filepath: str) -> bool:
        """
        Exports overall student attendance analytics and percentages to CSV.
        """
        rows = get_attendance_analytics()
        os.makedirs(os.path.dirname(os.path.abspath(filepath)), exist_ok=True)

        with open(filepath, mode="w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerow(["Roll Number", "Student Name", "Total Sessions", "Classes Present", "Attendance Percentage", "Status Flag"])
            for r in rows:
                roll, name, total, present, pct = r
                flag = "No attendance" if total == 0 else ("Eligible" if pct >= 75 else "ALERT: Low Attendance")
                writer.writerow([roll, name, total, present, f"{pct}%" if total else "", flag])

        return True

    @staticmethod
    def export_attendance_html_report(filepath: str, title: str = "EduPilot Attendance Summary",
                                      start_date: str = None, end_date: str = None,
                                      subject_id: int = None, section: str = None,
                                      class_id: int = None) -> bool:
        """
        Generates a printable HTML report document (convertible to PDF by browser/print).
        """
        rows = get_attendance_analytics(start_date, end_date, subject_id, section, class_id)
        os.makedirs(os.path.dirname(os.path.abspath(filepath)), exist_ok=True)

        if start_date and end_date and start_date == end_date:
            period_description = f"Date: {start_date}"
        elif start_date or end_date:
            period_description = f"Period: {start_date or 'Beginning'} – {end_date or 'Today'}"
        else:
            period_description = "Period: All available dates"
        if section:
            period_description += f" · Section {section}"

        table_rows_html = ""
        for r in rows:
            roll, name, total, present, pct = r
            badge_color = "#48677D" if total and pct >= 75 else "#7A6B52"
            status_text = "No sessions" if not total else ("Good" if pct >= 75 else "Low (<75%)")
            table_rows_html += f"""
            <tr>
                <td>{escape(str(roll))}</td>
                <td><strong>{escape(str(name))}</strong></td>
                <td>{total}</td>
                <td>{present}</td>
                <td><span style="color: {badge_color}; font-weight: bold;">{f'{pct}%' if total else '—'}</span></td>
                <td><span style="background-color: {badge_color}22; color: {badge_color}; padding: 4px 8px; border-radius: 4px; font-weight: 600;">{status_text}</span></td>
            </tr>
            """

        html_content = f"""<!DOCTYPE html>
<html>
<head>
    <meta charset="utf-8">
    <title>{escape(str(title))}</title>
    <style>
        @page {{ size: A4 landscape; margin: 12mm; }}
        body {{ font-family: 'Segoe UI', Arial, sans-serif; margin: 0; background-color: #FFFFFF; color: #26323B; font-size: 11pt; }}
        .header {{ border-bottom: 2pt solid #48677D; padding-bottom: 10pt; margin-bottom: 16pt; }}
        .header h1 {{ color: #263746; margin: 0; font-size: 20pt; }}
        .header p {{ color: #52616B; margin: 4pt 0 0 0; font-size: 11pt; }}
        table {{ width: 100%; border-collapse: collapse; margin-top: 10pt; background: white; }}
        thead {{ display: table-header-group; }}
        tr {{ break-inside: avoid; }}
        th {{ background-color: #26323B; color: white; text-align: left; padding: 8pt 9pt; font-size: 10pt; text-transform: uppercase; }}
        td {{ padding: 8pt 9pt; border-bottom: 1px solid #DCE2E7; font-size: 11pt; }}
        tr:nth-child(even) {{ background-color: #F4F6F8; }}
        .footer {{ margin-top: 20pt; font-size: 9pt; color: #52616B; text-align: center; border-top: 1px solid #DCE2E7; padding-top: 9pt; }}
    </style>
</head>
<body>
    <div class="header">
        <h1>{escape(str(title))}</h1>
        <p>Velammal College of Engineering &amp; Technology (VCET) — EduPilot System Report</p>
        <p>{escape(period_description)}</p>
    </div>

    <table>
        <thead>
            <tr>
                <th>Roll Number</th>
                <th>Student Name</th>
                <th>Total Classes</th>
                <th>Classes Attended</th>
                <th>Attendance %</th>
                <th>Status</th>
            </tr>
        </thead>
        <tbody>
            {table_rows_html}
        </tbody>
    </table>

    <div class="footer">
        Generated automatically by EduPilot Local Assistant • Velammal College of Engineering and Technology
    </div>
</body>
</html>"""

        with open(filepath, "w", encoding="utf-8") as f:
            f.write(html_content)

        return True

    @staticmethod
    def export_student_transcript_html(user_id: int, filepath: str) -> bool:
        """
        Generates an individual student academic & attendance transcript HTML report.
        """
        profile = get_student_profile(user_id)
        if not profile:
            return False

        att_summary = get_student_attendance(user_id)
        academic_rows = get_student_academic_records(user_id)
        os.makedirs(os.path.dirname(os.path.abspath(filepath)), exist_ok=True)

        rows_html = ""
        for code, name, total, present in att_summary:
            pct = round((present / total * 100), 1) if total > 0 else None
            color = "#48677D" if pct is not None and pct >= 75 else "#7A6B52"
            rows_html += f"""
            <tr>
                <td><strong>{escape(str(code))}</strong></td>
                <td>{escape(str(name))}</td>
                <td>{total}</td>
                <td>{present}</td>
                <td><span style="color: {color}; font-weight: bold;">{f'{pct}%' if pct is not None else '—'}</span></td>
            </tr>
            """

        academic_html = "".join(
            f"<tr><td>{escape(str(code))}</td><td>{escape(str(name))}</td>"
            f"<td>{internal:g}</td><td>{assignment:g}</td><td>{lab:g}</td><td>{total:g}/300</td></tr>"
            for code, name, internal, assignment, lab, total in academic_rows
        )

        html_content = f"""<!DOCTYPE html>
<html>
<head>
    <meta charset="utf-8">
    <title>Student Academic Transcript - {escape(str(profile[3]))}</title>
    <style>
        body {{ font-family: 'Segoe UI', Arial, sans-serif; margin: 40px; background-color: #ffffff; color: #26323B; }}
        .card {{ background: #FFFFFF; border: 1px solid #DCE2E7; border-radius: 10px; padding: 20px; margin-bottom: 25px; }}
        h1 {{ color: #263746; margin: 0 0 10px 0; font-size: 24px; }}
        .grid {{ display: grid; grid-template-columns: 1fr 1fr; gap: 15px; font-size: 14px; }}
        table {{ width: 100%; border-collapse: collapse; margin-top: 15px; }}
        th {{ background-color: #26323B; color: white; text-align: left; padding: 10px; font-size: 12px; }}
        td {{ padding: 10px; border-bottom: 1px solid #DCE2E7; font-size: 14px; }}
    </style>
</head>
<body>
    <div class="card">
        <h1>Student Performance Record</h1>
        <div class="grid">
            <div><strong>Name:</strong> {escape(str(profile[3]))}</div>
            <div><strong>Roll Number:</strong> {escape(str(profile[2]))}</div>
            <div><strong>Department:</strong> {escape(str(profile[4]))}</div>
            <div><strong>Year & Section:</strong> Year {profile[5]} - Section {profile[6]}</div>
        </div>
    </div>

    <h3>Published Academic Marks</h3>
    <table><thead><tr><th>Subject</th><th>Course</th><th>Internal /100</th><th>Assignment /100</th><th>Lab /100</th><th>Total</th></tr></thead>
    <tbody>{academic_html or '<tr><td colspan="6">No academic marks published</td></tr>'}</tbody></table>

    <h3>Subject Attendance</h3>
    <table>
        <thead>
            <tr>
                <th>Subject Code</th>
                <th>Subject Name</th>
                <th>Total Sessions</th>
                <th>Present</th>
                <th>Attendance %</th>
            </tr>
        </thead>
        <tbody>
            {rows_html}
        </tbody>
    </table>
</body>
</html>"""

        with open(filepath, "w", encoding="utf-8") as f:
            f.write(html_content)

        return True

report_generator = ReportGenerator()
